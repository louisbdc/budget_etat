"""Tests des parseurs sur des fichiers qui reproduisent les FORMATS observés dans
data/INSPECTION.md (en-têtes, séparateurs, virgules décimales, encodages).
Les montants sont abstraits : ce ne sont pas des données budgétaires.
"""

import io
import zipfile

import pytest

from budget_etat import ingest, inspect_raw, queries, sources
from budget_etat.db import connect
from budget_etat.fetch import select_datasets


def write(path, text, enc="utf-8"):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(text.encode(enc))
    return path


# --- utilitaires -----------------------------------------------------------------

@pytest.mark.parametrize("raw,expected", [
    ("1 234,56", 1234.56), ("-8836054,82", -8836054.82), ("1234.0", 1234.0), ("13 498", 13498.0),
    ("1.234,5", 1234.5), ("", None), (None, None), (42, 42.0),
])
def test_to_float(raw, expected):
    assert sources.to_float(raw) == expected


def test_split_programme_and_perimetre():
    assert sources.split_programme("Action de la France en Europe et dans le monde - 105") == \
        ("105", "Action de la France en Europe et dans le monde")
    assert sources.split_programme("Sans code") == (None, "Sans code")
    assert [sources.perimetre_programme(c) for c in ("105", "552", "613", "741", "862", "951", None)] == \
        ["BG", "BG", "BA", "CAS", "CCF", "CC", "inconnu"]


def test_year_from_path(tmp_path):
    assert sources.year_from_path(tmp_path / "economie/plrg-2024/attachments/annexe1_etat_cp_2024_csv") == 2024
    assert sources.year_from_path(tmp_path / "economie/execution-2011-du-budget-general-en-cp/export.csv") == 2011


def test_decode_cp1252():
    assert sources.decode_bytes("Prélèvement".encode("cp1252")) == ("Prélèvement", "cp1252")


def test_select_datasets():
    cat = [{"id": i, "title": t} for i, t in [
        ("plrg-2023", ""), ("plr-2019-donnees-execution", ""), ("execution-2016-du-budget-general-en-cp", ""),
        ("execution-2016-du-budget-general-en-ae", ""), ("plr-2015-projet-de-loi-de-reglement-pour-2015-donnees-du-volet-performance-prese", ""),
        ("plf-2026-budget-vert", "PLF 2026 - Budget vert"), ("x", "PLF 2019 - Recettes fiscales nettes"),
        ("y", "Annuaire statistique")]]
    assert select_datasets(cat) == ["plrg-2023", "plr-2019-donnees-execution",
                                    "execution-2016-du-budget-general-en-cp", "x"]


# --- formats d'exécution « par titre » ----------------------------------------------

EXEC_2011 = """mission;programme;libelle_programme;action;libelle_action;t1;t2;t3;t4;t5;t6;t7;total_general
Mission A;101;Programme A1;01;Action 1;;2000000;1000000;;;;;3000000
Mission A;101;Programme A1;02;Action 2;;500000;;;;;;500000
Engagements financiers de l'État;117;Charge de la dette et trésorerie de l'État;01;Dette;;;;4000000;;;;4000000
Remboursements et dégrèvements;200;Remboursements et dégrèvements d'impôts d'État;01;RD;;;;;;9000000;;9000000
Pensions;741;Pensions civiles;01;CAS;;7000000;;;;;;7000000
"""
EXEC_2010 = """ministere;libelle;pgm;libelle_pgm;action;libelle_action;t1;t2;t3;t4;t5;t6;t7;total_general
07;Ministère X;101;Programme A1;01;Action 1;;1000000;;;;;;1000000
"""
EXEC_2012 = """mission;programme;libelle_programme;action;libelle_action;t1;t2;t3;t4;t5;t6;t7
Mission A;101;Programme A1;01;Action 1;;3000000;;;;;
"""


@pytest.fixture
def raw(tmp_path):
    r = tmp_path / "raw"
    e = r / "economie"
    write(e / "execution-2011-du-budget-general-en-cp/export.csv", EXEC_2011)
    write(e / "execution-2010-du-budget-general-en-cp/export.csv", EXEC_2010)
    write(e / "execution-2012-du-budget-de-letat-en-cp-suivant-la-nomenclature-mission-programm/export.csv", EXEC_2012)
    write(e / "execution-2012-du-budget-de-letat-en-cp-suivant-la-nomenclature-ministere-progra/export.csv", EXEC_2012)
    write(e / "execution-2012-du-budget-de-letat-en-ae-suivant-la-nomenclature-mission-programm/export.csv", EXEC_2012)
    a = e / "plrg-2024" / "attachments"
    write(a / "annexe1_etat_titre_cat_2024_csv",
          "Mission;Programme;Titre;Categorie;Depenses\n"
          "Mission A;Programme A1 - 101;Titre 2;21;1500000,5\n"
          "Mission A;Programme A1 - 101;Titre 3;31;500000\n"
          "Engagements financiers de l'État;Charge de la dette et trésorerie de l'État - 117;Titre 4;41;4000000\n")
    write(a / "annexe1_etat_cp_2024_csv",
          "Mission;Programme;Titre;LFI;Depenses_constatees\n"
          "Mission A;Programme A1 - 101;Titre 2 - Dépenses de personnel;0;1500000,5\n"
          "Mission A;Programme A1 - 101;Autres titres - Autres dépenses;0;500000\n"
          "Engagements financiers de l'État;Charge de la dette et trésorerie de l'État - 117;Autres titres - Autres dépenses;0;4100000\n")
    write(a / "annexe1_etat_recettes_csv",
          "Niveau hiérarchique de la ligne;Catégorie;Section;Ligne de prévision;Ligne d'exécution;LFI;LFR/LFG;"
          "Total des prévisions;Total des recouvrements;Total des recettes**\n"
          "0;Recettes fiscales;11 - Impôt X;1101 - Impôt X;;0;0;0;0;0\n"
          "1;Recettes fiscales;11 - Impôt X;1101 - Impôt X;110101 - détail;0;0;0;7000000;0\n"
          "2;Recettes fiscales;11 - Impôt X (total);11 - Impôt X (total);;0;0;0;0;6000000\n"
          "2;Recettes fiscales;16 - Impôt Y (total);16 - Impôt Y (total);;0;0;0;0;2000000,5\n"
          "3;Recettes fiscales (total);Recettes fiscales (total);Recettes fiscales (total);;0;0;0;0;8000000,5\n"
          "2;Prélèvements sur les recettes de l'État;32 - Prélèvement UE;3201 - PSR UE;;0;0;0;0;-1000000\n"
          "3;Prélèvements sur les recettes de l'État (total);Prélèvements sur les recettes de l'État (total);x;;0;0;0;0;-1000000\n")
    write(a / "bacea_bilan_2024_csv", "a;b\n1;2\n")
    return r


@pytest.fixture
def db(raw, tmp_path):
    p = tmp_path / "b.duckdb"
    assert ingest.run(raw, p) == 0
    return connect(p, read_only=True)


def test_exec_aggregates_actions_and_keeps_bg_only(db):
    rows = db.execute("SELECT programme_code, titre_code, montant_meur FROM depense WHERE exercice = 2011 "
                      "ORDER BY 1, 2").fetchall()
    assert rows == [("101", "2", 2.5), ("101", "3", 1.0), ("117", "4", 4.0), ("200", "6", 9.0)]  # 741 exclu


def test_exec_2010_mission_deduced_from_nomenclature(db):
    assert db.execute("SELECT mission_lib, mission_code FROM depense WHERE exercice = 2010").fetchall() == \
        [("Mission A", "déduite")]


def test_exec_ministere_and_ae_files_skipped(db):
    assert db.execute("SELECT sum(montant_meur) FROM depense WHERE exercice = 2012").fetchone()[0] == 3.0


def test_plrg_titre_cat_and_controls(db, tmp_path):
    assert db.execute("SELECT sum(montant_meur) FROM depense WHERE exercice = 2024").fetchone()[0] == \
        pytest.approx(6.0000005)
    assert db.execute("SELECT categorie_code FROM depense WHERE exercice = 2024 AND titre_code = '2'").fetchone() == ("21",)
    report = (tmp_path / "VALIDATION.md").read_text()
    assert "programme 117" in report and "-0.10" in report  # écart signalé
    assert "programme 101" not in report  # contrôle exact, non listé


def test_plrg_recettes_section_totals(db, tmp_path):
    rows = dict(db.execute("SELECT poste_lib, montant_meur FROM recette WHERE exercice = 2024").fetchall())
    assert rows == {"Impôt X": 6.0, "Impôt Y": 2.0000005, "Prélèvement UE": 1.0}  # prélèvement stocké positif
    s = queries.series(db)["etat"]
    i = s["years"].index(2024)
    assert s["recettes_nettes"][i] == pytest.approx(7.0000005)
    assert s["depenses"][i] == pytest.approx(6.0000005)  # 200/201 exclus, aucun en 2024 ici
    assert "recettes fiscale" not in (tmp_path / "VALIDATION.md").read_text()  # 8,0000005 = 6 + 2,0000005


def test_net_view_excludes_remboursements(db):
    s = queries.series(db)["etat"]
    assert s["depenses"][s["years"].index(2011)] == pytest.approx(7.5)  # 9 M€ du programme 200 exclus
    assert s["charge_dette"][s["years"].index(2011)] == pytest.approx(4.0)


def test_unused_attachment_is_skipped_not_error(raw, tmp_path, capsys):
    ingest.run(raw, tmp_path / "c.duckdb")
    out = capsys.readouterr().out
    assert "0 erreur(s)" in out and "ignoré(s)" in out


def test_unknown_format_is_reported(tmp_path, capsys):
    raw = tmp_path / "raw"
    write(raw / "economie/plrg-2030/attachments/annexe1_etat_titre_cat_2030_csv", "foo;bar\n1;2\n")
    assert ingest.run(raw, tmp_path / "d.duckdb") == 0
    assert "en-têtes ['foo', 'bar']" in capsys.readouterr().out


# --- PLF ---------------------------------------------------------------------------

def test_plf_recettes_and_depenses(tmp_path):
    rec = write(tmp_path / "plf-2024-recettes-du-budget-general/export.csv",
                "annee;type_de_recettes;code_ligne_recettes;libelle;montant_recettes_plf;montant_recettes_lfi\n"
                "2024.0;Recettes fiscales;1101.0;Impôt X;1000000.0;2000000.0\n")
    out = list(sources.parse_plf_recettes(rec))
    assert [(r["nature"], r["poste_code"], r["montant_meur"]) for _, r in out] == [("plf", "1101", 1.0), ("lfi", "1101", 2.0)]
    dep = write(tmp_path / "plf25-depenses-2025-selon-destination/export.csv",
                "exercice;loi;typebudget;ministere;libelle_ministere;mission;libelle_mission;programme;libelle_programme;"
                "action;libelle_action;sous_action;libelle_sous_action;categorie;titre;autorisation_engagement;credit_de_paiement\n"
                "2025;PLF;BG;1;M;AA;Mission A;101;Programme A1;101-01;a;;;21;2;0;1000000.0\n"
                "2025;PLF;BG;1;M;AA;Mission A;101;Programme A1;101-02;b;;;21;2;0;500000.0\n"
                "2025;PLF;CAS;1;M;ZZ;CAS;741;P;741-01;c;;;21;2;0;9000000.0\n")
    out = [r for _, r in sources.parse_plf_depenses(dep)]
    assert len(out) == 1 and out[0]["montant_meur"] == 1.5 and out[0]["nature"] == "plf"


# --- INSEE SDMX ------------------------------------------------------------------------

SDMX = """<?xml version="1.0" encoding="UTF-8"?>
<message:StructureSpecificData xmlns:message="http://www.sdmx.org/resources/sdmxml/schemas/v2_1/message">
<message:DataSet>
<Series IDBANK="001739081" TITLE_FR="Encours de la dette négociable totale de l'État" UNIT_MULT="9" UNIT_MEASURE="EUROS">
  <Obs TIME_PERIOD="2023-12" OBS_VALUE="2,5"/><Obs TIME_PERIOD="2024-12" OBS_VALUE="3.0"/>
</Series>
<Series IDBANK="001711532" TITLE_FR="Encours de la dette négociable de l'État à court terme (maturité d'un an et moins) en euros" UNIT_MULT="9">
  <Obs TIME_PERIOD="2024-12" OBS_VALUE="1.0"/>
</Series>
</message:DataSet></message:StructureSpecificData>"""


def test_insee_dette(tmp_path):
    p = write(tmp_path / "insee/dette_negociable.xml", SDMX)
    rows = list(sources.parse_insee(p))
    dette = [r for t, r in rows if t == "agregat_etat"]
    assert [(r["exercice"], r["mois"], r["montant_meur"]) for r in dette] == [(2023, 12, 2500.0), (2024, 12, 3000.0)]
    assert sum(1 for t, _ in rows if t == "serie") == 3


def test_insee_without_unit_mult_is_refused(tmp_path):
    p = write(tmp_path / "x.xml", SDMX.replace(' UNIT_MULT="9" UNIT_MEASURE="EUROS"', ""))
    with pytest.raises(sources.UnknownFormat, match="UNIT_MULT"):
        list(sources.parse_insee(p))


# --- inspection -------------------------------------------------------------------------

def test_inspect_cp1252_zip_and_catalog(tmp_path):
    raw = tmp_path / "raw"
    write(raw / "a/f_csv", "Libellé;Montant\nPrélèvement;1 234\n", enc="cp1252")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("notice.txt", "Total des recettes** : nettes des remboursements")
        z.writestr("data.csv", "x;y\n1;2\n")
    (raw / "a/notices_zip").write_bytes(buf.getvalue())
    write(raw / "economie/_catalog_all.json", '[{"id": "plrg-2023", "title": "PLRG 2023"}]')
    write(raw / "insee/dette_negociable.xml", SDMX)
    rep = tmp_path / "I.md"
    assert inspect_raw.run(raw, rep) == 0
    t = rep.read_text()
    assert "encodage : cp1252" in t and "Prélèvement" in t
    assert "nettes des remboursements" in t and "(zip) `data.csv`" in t
    assert "`plrg-2023` | PLRG 2023" in t
    assert "001739081" in t and "2023-12 → 2024-12" in t


# --- formats ajoutés après la 2e inspection -------------------------------------------

PIVOT_2013 = """Gestion;2013;;;;;;;;;;;
Type Budget;All;;;;;;;;;;;
;;;;;;;;;;;;
CP;;;;;TITRE;;;;;;;
Mission;Programme;Libellé programme;Action;Libellé action;T1;T2;T3;T4;T5;T6;T7;Total général
Accords monétaires internationaux;811;Relations UMOA;01;x;;;;0;;;;0
Mission A;101;Programme A1;01;Action 1;;24 928 201;17 740 210;;98 139;-2 559;;42 764 991
Mission A;101;Programme A1;02;Action 2;;1 000 000;;;;;;1 000 000
Total général;;;;;;25 928 201;17 740 210;;98 139;-2 559;;43 764 991
"""


def test_exec_pivot_2013(tmp_path):
    p = write(tmp_path / "economie/execution-2013-du-budget-de-letat-en-cp-et-ae-/attachments/"
              "execution_2013_du_budget_de_l_etat_en_cp_suivant_la_nomenclature_mission_program", PIVOT_2013, "cp1252")
    rows = {(r["programme_code"], r["titre_code"]): r["montant_meur"] for _, r in sources.parse_exec_titres(p)}
    assert rows == {("101", "2"): pytest.approx(25.928201), ("101", "3"): pytest.approx(17.74021),
                    ("101", "5"): pytest.approx(0.098139), ("101", "6"): pytest.approx(-0.002559)}
    assert {r["exercice"] for _, r in sources.parse_exec_titres(p)} == {2013}


def test_exec_pivot_ae_and_ministere_skipped(tmp_path):
    for name in ("execution_2013_du_budget_de_l_etat_en_ae_et_cp_suivant_la_nomenclature_mission_p",
                 "execution_2013_du_budget_de_l_etat_en_cp_suivant_la_nomenclature_ministere_progr",
                 "plr2014_exec_min_cp_csv"):
        p = write(tmp_path / "economie/x-2013/attachments" / name, PIVOT_2013)
        with pytest.raises(sources.Skip):
            list(sources.parse_exec_titres(p))


DEST_NAT = ("annee_rap;type_de_budget_hors_budgets_annexes;code_mission;mission;code_programme;programme;code_action;"
            "action;code_sous_action;sous_action;code_categorie;categorie;exec_ae_2018_rap_2018;exec_cp_2018_rap_2018;"
            "code_titre;titre;code_ministere_au_1er_janvier_2018;ministere_au_1er_janvier_2018\n"
            "2018;BG;AA;Mission A;101;Programme A1;01;a;;;21;Rém;9;2000000;2;Personnel;MIN01;M\n"
            "2018;BG;AA;Mission A;101;Programme A1;02;b;;;21;Rém;9;1000000;2;Personnel;MIN01;M\n"
            "2018;BG;AA;Mission A;101;Programme A1;02;b;;;31;Fonct;9;500000;3;Fonct;MIN01;M\n"
            "2018;CAS;YE;CAS;753;P;01;c;;;31;F;9;7000000;3;F;MIN09;I\n")


def test_destination_nature_csv_and_controle(tmp_path):
    d = tmp_path / "raw/economie"
    write(d / "projet-de-loi-de-reglement-2019-plr-20190/export.csv", DEST_NAT, "utf-8-sig")
    write(d / "projet-de-loi-de-reglement-2019-plr-2019/export.csv",
          "annee_rap;type_de_budget;code_mission;mission;code_programme;programme;exec_t2_ae_cp_2018_rap_2018;"
          "exec_ae_hors_t2_2018_rap_2018;exec_cp_hors_t2_2018_rap_2018;exec_ae_t2_hors_t2_2018_rap_2018;"
          "exec_cp_t2_hors_t2_2018_rap_2018;exec_etpt_2018_rap_2018\n"
          "2018;Budget général;AA;Mission A;101;Programme A1;3000000;500000;500000;3500000;3500000;10\n"
          "2018;Comptes d'affectation spéciale;YE;CAS;753;P;;7000000;7000000;7000000;7000000;\n")
    write(d / "projet-de-loi-de-reglement-2019-plr-20191/export.csv",
          "annee_rap;type_de_budget_hors_budgets_annexes;code_programme;t1_exec_ae_2018;t1_exec_cp_2018\n2018;BG;101;;\n")
    dbp = tmp_path / "b.duckdb"
    assert ingest.run(tmp_path / "raw", dbp) == 0
    con = connect(dbp, read_only=True)
    assert con.execute("SELECT exercice, titre_code, categorie_code, montant_meur FROM depense ORDER BY 2").fetchall() == \
        [(2018, "2", "21", 3.0), (2018, "3", "31", 0.5)]  # exercice lu dans annee_rap (id du jeu : 2019)
    assert "programme 101" not in (tmp_path / "VALIDATION.md").read_text()  # contrôle exact


def test_destination_nature_xlsx(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Titre du tableau"])
    for line in DEST_NAT.strip().split("\n"):
        ws.append(line.split(";"))
    p = tmp_path / "economie/projet-de-loi-de-reglement-2020-plr-2020/attachments/plr2020_credits_destination_nature_xls"
    p.parent.mkdir(parents=True)
    wb.save(p)
    rows = [r for _, r in sources.parse_plr_attachment(p)]
    assert sum(r["montant_meur"] for r in rows) == pytest.approx(3.5)


RECETTES_NETTES_2014 = """(en million d'euros);column1;column2
Désignation des recettes;Exécution 2012;Évaluation initiale pour 2013
;;
 A. Recettes fiscales;358 997;394 780
 1 Impôt sur le revenu;65 510;77 298
 B. Remboursements et dégrèvements;90 559;96 163
 1. Impôt sur le revenu;6 030;5 396
3. Taxe sur la valeur ajoutée;51 262;54 500
 C. Recettes fiscales nettes;268 438;298 617
1. Impôt sur le revenu net (A.1 - B.1);59 480;71 902
2. Impôt sur les sociétés net (A.3 - B.2);40 832;53 531
3. TICPE (brute A5);13 498;13 680
4. Taxe sur la valeur ajoutée - nette (A.6 - B.3);133 403;141 245
5. Autres recettes fiscales - nettes (A.2 + A.3bis);21 224;18 259
 D. Recettes non fiscales;14 110;14 209
 E. Prélèvements sur les recettes de l'État;74 635;76 128
  Prélèvements sur les recettes de l'État au profit des collectivités territoriales;55 584;55 693
  Prélèvement sur les recettes de l’État au profit de l’Union européenne;19 052;20 435
"""


def test_recettes_nettes(tmp_path):
    p = write(tmp_path / "plf-2014-recettes-fiscales-nettes/attachments/plf_2014_recettes_fiscales_nettes_csv",
              RECETTES_NETTES_2014, "cp1252")
    out = list(sources.parse_recettes_nettes(p))
    rec = {r["poste_code"]: r["montant_meur"] for t, r in out if t == "recette"}
    assert rec == {"IR": 59480, "IS": 40832, "TICPE": 13498, "TVA": 133403, "AUTRES": 21224, "NF": 14110,
                   "PSR_COLL": 55584, "PSR_UE": 19052}
    assert {r["exercice"] for _, r in out} == {2012}
    ctl = {r["cle"]: r["reference"] for t, r in out if t == "controle"}
    assert ctl == {"recettes fiscale": 268438, "recettes prelevement": 74635}
    # Montants publiés arrondis au M€ : la somme des lignes diffère du total de 1 M€.
    assert sum(v for k, v in rec.items() if k in ("IR", "IS", "TICPE", "TVA", "AUTRES")) == pytest.approx(ctl["recettes fiscale"], abs=5)


def test_recettes_nettes_thousands_and_unknown_layout(tmp_path):
    p = write(tmp_path / "plf-2012/x_csv", "(En milliers d’euros);;Exécution 2010\n;Impôt net sur le revenu;47433070\n")
    with pytest.raises(sources.UnknownFormat, match="PLF 2012 incomplet"):
        list(sources.parse_recettes_nettes(p))


def test_inspect_xlsx_and_pdf(tmp_path):
    import openpyxl
    from pypdf import PdfWriter

    wb = openpyxl.Workbook()
    wb.active.append(["code_programme", "exec_cp_2019"])
    wb.active.append(["101", "12"])
    raw = tmp_path / "raw"
    (raw / "x").mkdir(parents=True)
    wb.save(raw / "x" / "plr2019_credits_destination_nature_xls")
    w = PdfWriter()
    w.add_blank_page(100, 100)
    buf = io.BytesIO()
    w.write(buf)
    (raw / "x" / "notice.pdf").write_bytes(buf.getvalue())
    rep = tmp_path / "I.md"
    assert inspect_raw.run(raw, rep) == 0
    t = rep.read_text()
    assert "feuille `Sheet`" in t and "code_programme ; exec_cp_2019" in t and "PDF, 1 pages" in t


# --- SMB (situations mensuelles budgétaires) ---------------------------------------------

SMB_HEAD = ("niveau_hierarchique;niveau_hierarchique_de_la_ligne;categorie;sous_categorie;ligne_d_information;"
            "30_11_{y};31_12_{y}\n")
SMB_ROWS = """0;Nul;Solde budgétaire;Solde budgétaire;Solde budgétaire;-150000000;-160000000
1;Sous-total de niveau 1;Dépenses;Budget général;Total dépenses nettes du budget général;9000000;10000000
2;Sous-total de niveau 2;Dépenses;Budget général;Charges de la dette de l’Etat;1000000;1500000
2;Sous-total de niveau 2;Dépenses;Prélèvements sur recettes;PSR au profit de l'Union européenne;100000;200000
3;Sous-total de niveau 3;Recettes;Budget général;Impôt sur le revenu;3000000;4000000
3;Sous-total de niveau 3;Recettes;Budget général;Taxe sur la valeur ajoutée;2000000;3000000
2;Sous-total de niveau 2;Recettes;Budget général;Total recettes non fiscales;500000;1000000
4;Sous-total de niveau 4;Soldes;Comptes spéciaux;CCF Avances aux collectivités territoriales;1;2
"""


def test_smb_parser(tmp_path):
    p = write(tmp_path / "export.csv", SMB_HEAD.format(y=2024) + SMB_ROWS, "utf-8-sig")
    out = list(sources.parse_smb(p))
    agg = {(r["indicateur"], r["mois"]): r["montant_meur"] for t, r in out if t == "agregat_etat"}
    assert agg[("solde", 12)] == -160 and agg[("charge_dette", 12)] == 1.5 and agg[("depenses_nettes", 11)] == 9
    rec = {(r["poste_code"], r["mois"]): r["montant_meur"] for t, r in out if t == "recette"}
    assert rec == {("PSR_UE", 11): 0.1, ("PSR_UE", 12): 0.2, ("IR", 11): 3, ("IR", 12): 4, ("TVA", 11): 2,
                   ("TVA", 12): 3, ("NF", 11): 0.5, ("NF", 12): 1}
    assert sum(1 for t, _ in out if t == "serie") == 16  # toutes les lignes gardées en série brute


def test_smb_priority_dedup_and_official_series(tmp_path):
    raw = tmp_path / "raw"
    d = raw / "economie/situations-mensuelles-budgetaires-series-longues"
    write(d / "export.csv", SMB_HEAD.format(y=2024) + SMB_ROWS, "utf-8-sig")
    write(d / "attachments/serie_longue_smb_dgfip_2024_xxcsv", SMB_HEAD.format(y=2024) + SMB_ROWS)
    a = raw / "economie/plrg-2024/attachments"
    write(a / "annexe1_etat_titre_cat_2024_csv", "Mission;Programme;Titre;Categorie;Depenses\n"
          "Mission A;Programme A1 - 101;Titre 2;21;8000000\n"
          "Engagements financiers de l'État;Charge - 117;Titre 4;41;1500000\n"
          "Remboursements et dégrèvements;R&D locaux - 201;Titre 2;21;500000\n"
          "Remboursements et dégrèvements;R&D État - 200;Titre 2;21;99000000\n")
    write(a / "annexe1_etat_recettes_csv",
          "Niveau hiérarchique de la ligne;Catégorie;Section;Ligne de prévision;Ligne d'exécution;LFI;LFR/LFG;"
          "Total des prévisions;Total des recouvrements;Total des recettes**\n"
          "2;Recettes fiscales;11 - Impôt X (total);x;;0;0;0;0;7000000\n")
    dbp = tmp_path / "b.duckdb"
    assert ingest.run(raw, dbp) == 0
    con = connect(dbp, read_only=True)
    # SMB prioritaire sur PLRG pour les recettes ; doublon export / pièce jointe éliminé
    srcs = con.execute("SELECT DISTINCT source FROM recette").fetchall()
    assert len(srcs) == 1 and srcs[0][0].startswith("SMB")
    assert con.execute("SELECT count(*) FROM agregat_etat WHERE indicateur = 'solde' AND mois = 12").fetchone()[0] == 1
    s = queries.series(con)["etat"]
    i = s["years"].index(2024)
    assert s["solde"][i] == -160 and "officiel" in s["solde_source"][i]
    assert s["depenses"][i] == 10  # agrégat SMB
    assert s["recettes_nettes"][i] == pytest.approx(4 + 3 + 1 - 0.2)
    report = (tmp_path / "VALIDATION.md").read_text()
    # détail : 8 + 1,5 (P117) + 0,5 (P201 gardé depuis 2023) = 10 = SMB ; P200 exclu
    assert "| 2024 | dépenses nettes | 10 | 10 | 0.0 |" in report
    assert "PLRG 2024 annexe1_etat_recettes" in report  # source écartée, journalisée


# --- 4e inspection : fins de ligne CR, PLR 2019-2020 (xls + nomenclature) -------------------

def test_cr_only_line_endings(tmp_path):
    p = tmp_path / "series_longues_smb_dgfip_2013_2023_csv"
    p.write_bytes((SMB_HEAD.format(y=2013) + SMB_ROWS).replace("\n", "\r").encode("utf-8"))
    agg = {(r["indicateur"], r["exercice"], r["mois"]): r["montant_meur"]
           for t, r in sources.parse_smb(p) if t == "agregat_etat"}
    assert agg[("solde", 2013, 12)] == -160


def _xlsx(path, rows):
    import openpyxl

    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def test_destination_nature_plr2019_variant_with_nomenclature(tmp_path):
    d = tmp_path / "economie/projet-de-loi-de-reglement-2019-plr-20192/attachments"
    _xlsx(d / "plr2019_credits_destination_nature_xls", [
        ["exercice", "loi", "typeBudget", "ministere", "mission", "programme", "action", "sous_action",
         "categorie", "titre", "AE EXEC", "CP EXEC"],
        ["2019", "PLR", "BG", "01", "AA", "105", "105-01", None, "22.0", "2", 1.0, 2000000.0],
        ["2019", "PLR", "BG", "01", "AA", "105", "105-02", None, "22.0", "2", 1.0, 1000000.0],
        [2020.0, "PLR", "BG", 7.0, "TR", 348.0, "348-11", None, 31.0, 3.0, 1.0, 500000.0],
        ["2019", "PLR", "CAS", "21", "YK", "793", "793-08", None, "31", "3", 1.0, 9000000.0],
    ])
    _xlsx(d / "plr2019_nomenclature_xls", [
        ["Type ligne", "Type Budget", "code", "Mission", "Ministere", "Libelle", "Libelle abrege", "commentFP"],
        ["MSN", "BG", "AA", None, None, "Action extérieure de l'État", "AEE", None],
        ["PGM", "BG", "105.0", "AA", "01", "Action de la France en Europe et dans le monde", "x", None],
        ["ACT", "BG", "105-01", None, None, "Coordination", None, None],
    ])
    rows = [r for _, r in sources.parse_plr_attachment(d / "plr2019_credits_destination_nature_xls")]
    got = {(r["exercice"], r["programme_code"], r["titre_code"], r["categorie_code"]): r for r in rows}
    assert set(got) == {(2019, "105", "2", "22"), (2020, "348", "3", "31")}  # CAS exclu, actions agrégées
    r = got[(2019, "105", "2", "22")]
    assert r["montant_meur"] == 3.0 and r["mission_lib"] == "Action extérieure de l'État"
    assert r["programme_lib"] == "Action de la France en Europe et dans le monde"
    assert got[(2020, "348", "3", "31")]["mission_lib"] == "TR"  # code gardé faute de libellé
    with pytest.raises(sources.Skip):  # le classeur de nomenclature n'est pas une source de montants
        list(sources.parse_plr_attachment(d / "plr2019_nomenclature_xls"))


def test_smb_date_formats():
    assert [sources._smb_date(h) for h in ("31_12_2024", "31/12/2013", "2013-12-31", "2013-12", "12/2013",
                                            "2013-12-31 00:00:00", "niveau", "13_2024")] == \
        [(2024, 12), (2013, 12), (2013, 12), (2013, 12), (2013, 12), (2013, 12), None, None]


def test_validation_lists_file_status(tmp_path):
    raw = tmp_path / "raw"
    write(raw / "economie/plrg-2030/attachments/annexe1_etat_titre_cat_2030_csv", "foo;bar\n1;2\n")
    write(raw / "economie/plrg-2030/attachments/bacea_bilan_2030_csv", "a;b\n1;2\n")
    ingest.run(raw, tmp_path / "d.duckdb")
    t = (tmp_path / "VALIDATION.md").read_text()
    assert "## Statut des fichiers" in t and "| à inspecter | `economie/plrg-2030/attachments/annexe1_etat_titre_cat_2030_csv`" in t
    assert "bacea_bilan" not in t  # les fichiers ignorés sont seulement comptés


# --- 5e retour : apostrophe typographique, exercice manquant, formats ignorés ------------------

def test_smb_typographic_apostrophe_header(tmp_path):
    head = "Niveau hiérarchique;Catégorie;Sous-catégorie;Ligne d’information;30/11/2013;31/12/2013\n"
    rows = "\n".join(";".join(l.split(";")[0:1] + l.split(";")[2:]) for l in SMB_ROWS.strip().split("\n")) + "\n"
    p = tmp_path / "series_longues_smb_dgfip_2013_2023_csv"
    p.write_bytes((head + rows).replace("\n", "\r").encode("cp1252"))
    agg = {(r["indicateur"], r["exercice"], r["mois"]): r["montant_meur"]
           for t, r in sources.parse_smb(p) if t == "agregat_etat"}
    assert agg[("solde", 2013, 12)] == -160 and agg[("depenses_nettes", 2013, 11)] == 9


def test_plr2020_rows_without_exercice(tmp_path):
    d = tmp_path / "economie/projet-de-loi-de-reglement-2020-plr-2020/attachments"
    _xlsx(d / "plr2020_credits_destination_nature_xls", [
        ["exercice", "loi", "typeBudget", "ministere", "mission", "programme", "action", "sous_action",
         "categorie", "titre", "AE EXEC", "CP EXEC"],
        [2020.0, "PLR", "BG", 7.0, "TR", 348.0, "348-11", None, 31.0, 3.0, 1.0, 500000.0],
        [None, "PLR", "BG", 7.0, "TR", 348.0, "348-12", None, 31.0, 3.0, 1.0, 250000.0],
    ])
    rows = [r for _, r in sources.parse_plr_attachment(d / "plr2020_credits_destination_nature_xls")]
    assert [(r["exercice"], r["montant_meur"]) for r in rows] == [(2020, 0.75)]


def test_plrg_gross_recettes_and_empty_exports_are_skipped(tmp_path):
    p = write(tmp_path / "plrg-2023/attachments/annexe1_etat_recettes_2023_csv",
              "Categorie;Section;Ligne_prevision;Ligne_d'execution;LFI;LFR;Total_prevision;Total_recouvrement\n"
              "Recettes fiscales;11 - IR;1101 - IR;110101 - x;0;0;0;12\n")
    with pytest.raises(sources.Skip, match="SMB"):
        list(sources.parse_plr_attachment(p))
    vide = "recordid;_record_id;record_timestamp;_record_timestamp;record_size;_record_size;resource_id;_resource_id\n"
    for path, fn in [(tmp_path / "projet-de-loi-de-reglement-2020-plr-2020/export.csv", sources.parse_plr_export),
                     (tmp_path / "execution-2013-du-budget-de-letat-en-cp-et-ae-/export.csv", sources.parse_exec_titres)]:
        write(path, vide)
        with pytest.raises(sources.Skip, match="pièces jointes"):
            list(fn(path))


# --- 6e retour : UTF-16, cellule parasite « Expr2 » -----------------------------------------

def test_smb_utf16(tmp_path):
    head = "Niveau hiérarchique;Catégorie;Sous-catégorie;Ligne d’information;30/11/2013;31/12/2013\r\n"
    rows = "\r\n".join(";".join(l.split(";")[0:1] + l.split(";")[2:]) for l in SMB_ROWS.strip().split("\n"))
    p = tmp_path / "series_longues_smb_dgfip_2013_2023_csv"
    p.write_bytes((head + rows).encode("utf-16"))  # avec BOM FF FE
    assert sources.decode_bytes(p.read_bytes())[1] == "utf-16"
    agg = {(r["indicateur"], r["mois"]): r["montant_meur"] for t, r in sources.parse_smb(p) if t == "agregat_etat"}
    assert agg[("solde", 12)] == -160
    q = tmp_path / "sans_bom"
    q.write_bytes((head + rows).encode("utf-16-le"))
    assert sources.decode_bytes(q.read_bytes())[1] == "utf-16"


def test_plr_amount_garbage_is_skipped_and_reported(tmp_path):
    d = tmp_path / "raw/economie/projet-de-loi-de-reglement-2020-plr-2020/attachments"
    head = ["exercice", "loi", "typeBudget", "ministere", "mission", "programme", "action", "sous_action",
            "categorie", "titre", "AE EXEC", "CP EXEC"]
    ok = [[2020.0, "PLR", "BG", 7.0, "TR", 348.0, f"348-{i:02d}", None, 31.0, 3.0, 1.0, 1000000.0] for i in range(150)]
    _xlsx(d / "plr2020_credits_destination_nature_xls",
          [head] + ok + [[2020.0, "PLR", "BG", 7.0, "TR", 348.0, "348-99", None, 31.0, 3.0, "Expr1", "Expr2"]])
    assert ingest.run(tmp_path / "raw", tmp_path / "b.duckdb") == 0
    con = connect(tmp_path / "b.duckdb", read_only=True)
    assert con.execute("SELECT sum(montant_meur) FROM depense").fetchone()[0] == 150
    t = (tmp_path / "VALIDATION.md").read_text()
    assert "| ingéré ⚠ |" in t and "Expr2" in t


def test_destination_nature_keeps_only_plr_rows(tmp_path):
    d = tmp_path / "raw/economie/projet-de-loi-de-reglement-2020-plr-2020/attachments"
    head = ["exercice", "loi", "typeBudget", "ministere", "mission", "programme", "action", "sous_action",
            "categorie", "titre", "AE EXEC", "CP EXEC"]
    rows = [[2020.0, loi, "BG", 7.0, "TR", 348.0, "348-11", None, 31.0, 3.0, 1.0, v]
            for loi, v in [("PLR", 1000000.0), ("LFI", 900000.0), ("LFR", 950000.0)]]
    _xlsx(d / "plr2020_credits_destination_nature_xls", [head] + rows)
    assert ingest.run(tmp_path / "raw", tmp_path / "b.duckdb") == 0
    con = connect(tmp_path / "b.duckdb", read_only=True)
    assert con.execute("SELECT sum(montant_meur) FROM depense").fetchone()[0] == 1.0
    assert "lignes des lois ['LFI', 'LFR'] écartées" in (tmp_path / "VALIDATION.md").read_text()


def test_diagnostic_for_suspect_year(tmp_path):
    raw = tmp_path / "raw"
    d = raw / "economie/situations-mensuelles-budgetaires-series-longues"
    write(d / "export.csv", SMB_HEAD.format(y=2024) + SMB_ROWS, "utf-8-sig")  # SMB : dépenses nettes 10, T4 1,5
    a = raw / "economie/plrg-2024/attachments"
    write(a / "annexe1_etat_titre_cat_2024_csv", "Mission;Programme;Titre;Categorie;Depenses\n"
          "Mission A;Programme A1 - 101;Titre 6;61;30000000\n"  # gonflé
          "Engagements financiers de l'État;Charge - 117;Titre 4;41;1500000\n")
    ingest.run(raw, tmp_path / "b.duckdb")
    t = (tmp_path / "VALIDATION.md").read_text()
    assert "| 2024 ⚠ | dépenses nettes |" in t and "### Diagnostic 2024" in t
    assert "| 4 | 2 | 2 | 0 | 1.00 |" in t  # titre 4 exact (1,5 M€ arrondi)
    assert "| 101 Programme A1 |" in t
